# -*- coding: utf-8 -*-
"""
export_entidades_excluidas.py
==============================
Genera/actualiza, para cada entidad EXCLUIDA de la proyeccion futura por un
evento de intervencion/disolucion SBS confirmado (ver eventos_intervencion.py
y dataset_builder.excluir_posteriores_a_intervencion), dos artefactos:

  1. Historial de clasificaciones por clasificadora (ECR), con estructura
     analoga a data/clasificaciones/<CORTE>.xls (columnas "Tipo de Entidad",
     "Entidad" + una columna por ECR) pero con una fila por corte semestral
     (el historial completo), no una fila por archivo.
  2. Ficha de evento: resumen de por que se excluyo la entidad (tipo de
     evento, fecha, resolucion SBS, fuente, conteos).

Organizacion en disco (ver config.EXCLUIDAS_DIR), UNA CARPETA POR TIPO DE
ENTIDAD y, dentro, UN SOLO LIBRO EXCEL POR TIPO DE ARTEFACTO (no uno por
entidad): cada entidad es una HOJA dentro de ese libro.

    outputs/analisis/entidades_excluidas/
    ├── CMAC/
    │   ├── historial_clasificaciones_CMAC.xlsx   (hoja "CMAC Sullana", ...)
    │   └── ficha_evento_CMAC.xlsx                 (hoja "CMAC Sullana", ...)
    ├── CRAC/
    │   ├── historial_clasificaciones_CRAC.xlsx
    │   └── ficha_evento_CRAC.xlsx
    └── FINANCIERA/
        ├── historial_clasificaciones_FINANCIERA.xlsx  (hojas: cada financiera)
        └── ficha_evento_FINANCIERA.xlsx

Reglas de actualizacion (IDEMPOTENCIA -- este modulo se re-ejecuta cada vez
que se corre el pipeline, p.ej. al agregarse una nueva entidad quebrada en
un corte futuro):
  - Carpeta de destino (por tipo de entidad): si existe, se reutiliza tal
    cual (no se borra su contenido); si no existe, se crea.
  - Libro Excel (por tipo de artefacto x tipo de entidad): si el archivo ya
    existe en disco, se ABRE y se conserva; si no existe, se crea uno nuevo.
  - Hoja (por entidad): si la hoja de esa entidad ya existe dentro del
    libro, se REEMPLAZA su contenido (se recrea con los datos actuales,
    para reflejar nuevos cortes de clasificacion); si no existe, se agrega.
    Las hojas de OTRAS entidades del mismo tipo, ya presentes en el libro,
    nunca se tocan.
"""

from __future__ import annotations
import os
import re
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .. import config
from ..utils import log, section, normalize_entity_name
from . import sbs_loader, financial_loader, entity_homology, eventos_intervencion
from .ratings import normalize_rating

# ---------------------------------------------------------------------------
# Estilos (compartidos por ambas hojas)
# ---------------------------------------------------------------------------
_HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
_TITLE_FONT = Font(name="Arial", bold=True, size=13, color="1F4E78")
_SUBTITLE_FONT = Font(name="Arial", italic=True, size=10, color="595959")
_NOTE_FONT = Font(name="Arial", italic=True, size=9, color="833C00")
_NORMAL_FONT = Font(name="Arial", size=10)
_BOLD_FONT = Font(name="Arial", bold=True, size=10)
_EVENT_FILL = PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


# ---------------------------------------------------------------------------
# sheet_name_for_entity
# ---------------------------------------------------------------------------
# Que hace: convierte un nombre de entidad en un nombre de hoja Excel valido
# (<=31 caracteres, sin los caracteres prohibidos por el formato : \ / ? * [ ]).
# Por que asi: cada entidad excluida es una hoja fija dentro de un libro
# compartido por tipo de entidad; el nombre debe ser estable entre
# ejecuciones para que la logica de "actualizar si existe" funcione.
# ---------------------------------------------------------------------------
def sheet_name_for_entity(entity_name: str) -> str:
    name = re.sub(r"[:\\/?*\[\]]", "-", str(entity_name)).strip()
    return name[:31]


def _autosize(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _replace_sheet(wb: Workbook, sheet_name: str) -> Worksheet:
    """Si la hoja ya existe la elimina y la vuelve a crear vacia (misma
    posicion relativa no garantizada, pero mismo nombre); si no existe, la
    crea. Nunca toca otras hojas del libro."""
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(title=sheet_name)
    # Un libro nuevo de openpyxl trae una hoja "Sheet" por defecto y vacia;
    # se elimina una vez que ya hay al menos una hoja real con datos.
    if "Sheet" in wb.sheetnames and wb["Sheet"].max_row == 1 and wb["Sheet"].max_column == 1 \
            and wb["Sheet"]["A1"].value is None and len(wb.sheetnames) > 1:
        del wb["Sheet"]
    return ws


def _open_or_create_workbook(path: str) -> Workbook:
    if os.path.exists(path):
        log(f"  Libro existente encontrado, se actualiza: {path}")
        return load_workbook(path)
    log(f"  Libro no existe, se crea: {path}")
    return Workbook()


# ---------------------------------------------------------------------------
# entidades_excluidas_con_metadata
# ---------------------------------------------------------------------------
# Que hace: resuelve, para cada fila de eventos_intervencion.tabla_eventos(),
# su ENTITY_ID/ENTITY_NAME/ENTITY_TYPE en el universo canonico (aplicando la
# homologacion ya existente en entity_homology.py / config.py -- ver docstring
# de ese modulo para el detalle de por que NO hace falta agregar nada nuevo
# alli para estos 4 casos).
# Recibe: registry (EntityRegistry ya construido con build_from_financial_sources).
# Devuelve: lista de dicts {entity_id, entity_name, entity_type, evento(Series)}.
# ---------------------------------------------------------------------------
def entidades_excluidas_con_metadata(registry: entity_homology.EntityRegistry) -> list[dict]:
    eventos = eventos_intervencion.tabla_eventos()
    name_to_id = dict(zip(registry.canonical_df["ENTITY_NAME_NORM"], registry.canonical_df["ENTITY_ID"]))
    name_to_name = dict(zip(registry.canonical_df["ENTITY_NAME_NORM"], registry.canonical_df["ENTITY_NAME"]))
    name_to_type = dict(zip(registry.canonical_df["ENTITY_NAME_NORM"], registry.canonical_df["ENTITY_TYPE"]))

    resultado = []
    for _, ev in eventos.iterrows():
        norm = ev["ENTITY_NAME_NORM"]
        entity_id = name_to_id.get(norm)
        if entity_id is None:
            log(f"ADVERTENCIA: evento de intervencion para '{norm}' no homologa a ninguna "
                f"entidad del universo financiero; se omite de la exportacion.", level="WARN")
            continue
        resultado.append({
            "entity_id": entity_id,
            "entity_name": name_to_name[norm],
            "entity_type": name_to_type[norm],
            "evento": ev,
        })
    return resultado


# ---------------------------------------------------------------------------
# _historial_clasificaciones_entidad
# ---------------------------------------------------------------------------
# Que hace: arma la tabla de historial ECR de una entidad (una fila por
# corte semestral), con la MISMA estructura de columnas que los archivos
# crudos de data/clasificaciones (Tipo de Entidad, Entidad + una columna por
# ECR), agregando Corte y Fecha de corte para poder leer la evolucion.
# ---------------------------------------------------------------------------
def _historial_clasificaciones_entidad(sbs_long_resolved: pd.DataFrame, entity_id: str) -> pd.DataFrame:
    sub = sbs_long_resolved[sbs_long_resolved["ENTITY_ID"] == entity_id].copy()
    cortes_orden = list(config.CORTES_CLASIFICACION.keys())
    sub["CORTE_ORDER"] = sub["CORTE"].apply(lambda c: cortes_orden.index(c) if c in cortes_orden else -1)

    pivot = sub.pivot_table(
        index=["CORTE", "CORTE_ORDER", "FECHA_CORTE", "ENTITY_TYPE_RAW", "ENTITY_NAME_RAW"],
        columns="ECR", values="RATING_RAW", aggfunc="first",
    ).reset_index()
    pivot = pivot.sort_values("CORTE_ORDER").drop(columns=["CORTE_ORDER"])
    for ecr in config.ECR_LIST:
        if ecr not in pivot.columns:
            pivot[ecr] = np.nan
    pivot = pivot.rename(columns={
        "ENTITY_TYPE_RAW": "Tipo de Entidad", "ENTITY_NAME_RAW": "Entidad",
        "CORTE": "Corte", "FECHA_CORTE": "Fecha de corte",
    })
    return pivot[["Corte", "Fecha de corte", "Tipo de Entidad", "Entidad"] + config.ECR_LIST]


# ---------------------------------------------------------------------------
# _escribir_hoja_historial
# ---------------------------------------------------------------------------
def _escribir_hoja_historial(ws: Worksheet, entity_name: str, historial: pd.DataFrame, evento: pd.Series) -> None:
    cols_final = list(historial.columns)
    fecha_evento = pd.to_datetime(evento["FECHA_EVENTO"])

    ws["A1"] = f"Historial de clasificaciones por clasificadora (ECR) — {entity_name}"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols_final))
    ws["A2"] = ("Entidad excluida de la proyección futura (evento de intervención/disolución SBS "
                "confirmado). Estructura análoga a data/clasificaciones/<CORTE>.xls, con una fila "
                "por corte semestral (historial completo) en vez de una fila por entidad.")
    ws["A2"].font = _SUBTITLE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols_final))
    ws["A3"] = (f"Evento: {evento['TIPO_EVENTO']} el {fecha_evento.date()} — {evento['RESOLUCION_SBS']}. "
                f"Fuente: {evento['FUENTE']}")
    ws["A3"].font = _SUBTITLE_FONT
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(cols_final))

    header_row = 5
    for j, col in enumerate(cols_final, start=1):
        c = ws.cell(row=header_row, column=j, value=col)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _BORDER

    r = header_row + 1
    for _, row in historial.iterrows():
        fecha_corte_dt = pd.to_datetime(row["Fecha de corte"])
        es_posterior_evento = fecha_corte_dt >= fecha_evento
        for j, col in enumerate(cols_final, start=1):
            val = row[col]
            if pd.isna(val):
                val = ""
            c = ws.cell(row=r, column=j, value=val)
            c.font = _NORMAL_FONT
            c.border = _BORDER
            if es_posterior_evento:
                c.fill = _EVENT_FILL
        r += 1

    note_row = r + 1
    ws.cell(row=note_row, column=1,
            value="Filas resaltadas: cortes publicados en la misma fecha o después del evento de "
                  "intervención (la entidad siguió reportando/clasificándose durante el proceso de "
                  "liquidación/traspaso; no se usan para generar predicciones futuras, ver "
                  "motor/datos/eventos_intervencion.py).")
    ws.cell(row=note_row, column=1).font = _NOTE_FONT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(cols_final))

    _autosize(ws, [10, 13, 24, 22, 9, 9, 9, 11, 9, 9][:len(cols_final)])
    ws.freeze_panes = "A6"


# ---------------------------------------------------------------------------
# _escribir_hoja_ficha
# ---------------------------------------------------------------------------
def _escribir_hoja_ficha(ws: Worksheet, entity_name: str, entity_id: str, entity_type: str,
                          evento: pd.Series, n_cortes: int, ultimo_corte: str) -> None:
    fecha_evento = pd.to_datetime(evento["FECHA_EVENTO"])

    ws["A1"] = f"Ficha de exclusión — {entity_name}"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:B1")

    campos = [
        ("Entidad (nombre canónico)", entity_name),
        ("ENTITY_ID", entity_id),
        ("Tipo de entidad", entity_type),
        ("Motivo de exclusión", "Evento de intervención/disolución SBS confirmado antes o durante "
                                 "la ventana de predicción"),
        ("Tipo de evento", evento["TIPO_EVENTO"]),
        ("Fecha del evento", str(fecha_evento.date())),
        ("Resolución SBS", evento["RESOLUCION_SBS"]),
        ("Descripción", evento["DESCRIPCION"]),
        ("Fuente", evento["FUENTE"]),
        ("Último corte de clasificación disponible", ultimo_corte),
        ("N° de cortes de clasificación con datos", str(n_cortes)),
    ]
    r = 3
    for label, value in campos:
        ws.cell(row=r, column=1, value=label).font = _BOLD_FONT
        ws.cell(row=r, column=2, value=value).font = _NORMAL_FONT
        ws.cell(row=r, column=1).border = _BORDER
        ws.cell(row=r, column=2).border = _BORDER
        r += 1
    _autosize(ws, [34, 90])


# ---------------------------------------------------------------------------
# exportar_entidades_excluidas
# ---------------------------------------------------------------------------
# Que hace: orquesta la generacion/actualizacion de los 2 artefactos
# (historial de clasificaciones + ficha de evento) para todas las entidades
# excluidas, agrupadas por tipo de entidad, aplicando las reglas de
# actualizacion idempotente descritas en el docstring del modulo.
# Recibe: sbs_long_resolved (DataFrame largo de sbs_loader con columna
# ENTITY_ID ya resuelta via entity_homology), registry (EntityRegistry).
# Devuelve: None (efecto: escribe/actualiza archivos en config.EXCLUIDAS_DIR).
# ---------------------------------------------------------------------------
def exportar_entidades_excluidas(sbs_long_resolved: pd.DataFrame,
                                  registry: entity_homology.EntityRegistry) -> None:
    section("EXPORTACION: ficha de evento + historial de clasificaciones (entidades excluidas)")

    entidades = entidades_excluidas_con_metadata(registry)
    if not entidades:
        log("No hay entidades excluidas por intervencion/disolucion; no se genera nada.", level="WARN")
        return

    por_tipo: dict[str, list[dict]] = {}
    for ent in entidades:
        por_tipo.setdefault(ent["entity_type"], []).append(ent)

    for tipo, lista_entidades in por_tipo.items():
        carpeta = os.path.join(config.EXCLUIDAS_DIR, tipo)
        if os.path.isdir(carpeta):
            log(f"Carpeta de destino ya existe, se reutiliza: {carpeta}")
        else:
            log(f"Carpeta de destino no existe, se crea: {carpeta}")
        os.makedirs(carpeta, exist_ok=True)

        path_historial = os.path.join(carpeta, f"historial_clasificaciones_{tipo}.xlsx")
        path_ficha = os.path.join(carpeta, f"ficha_evento_{tipo}.xlsx")
        wb_historial = _open_or_create_workbook(path_historial)
        wb_ficha = _open_or_create_workbook(path_ficha)

        for ent in lista_entidades:
            entity_name = ent["entity_name"]
            sheet_name = sheet_name_for_entity(entity_name)
            historial = _historial_clasificaciones_entidad(sbs_long_resolved, ent["entity_id"])

            log(f"  [{tipo}] {entity_name}: {len(historial)} cortes -> hoja '{sheet_name}' "
                f"en {os.path.basename(path_historial)} y {os.path.basename(path_ficha)}")

            ws_h = _replace_sheet(wb_historial, sheet_name)
            _escribir_hoja_historial(ws_h, entity_name, historial, ent["evento"])

            ws_f = _replace_sheet(wb_ficha, sheet_name)
            ultimo_corte = historial["Corte"].iloc[-1] if len(historial) else "N/A"
            _escribir_hoja_ficha(ws_f, entity_name, ent["entity_id"], tipo, ent["evento"],
                                  len(historial), ultimo_corte)

        wb_historial.save(path_historial)
        wb_ficha.save(path_ficha)
        log(f"  Guardado: {path_historial}")
        log(f"  Guardado: {path_ficha}")


# ---------------------------------------------------------------------------
# construir_insumos
# ---------------------------------------------------------------------------
# Que hace: reconstruye, a partir de las fuentes crudas, los dos insumos que
# necesita exportar_entidades_excluidas (universo canonico homologado +
# clasificaciones SBS con ENTITY_ID resuelto). Se deja como funcion separada
# para que scripts/exportar_entidades_excluidas.py sea un wrapper minimo, y
# para poder reutilizar el registry/sbs_long ya construidos por otros
# scripts (run_evaluacion.py, run_prediccion.py) en vez de recalcularlos.
# ---------------------------------------------------------------------------
def construir_insumos() -> tuple[pd.DataFrame, entity_homology.EntityRegistry]:
    sources = financial_loader.load_all_financial_sources()
    registry = entity_homology.EntityRegistry()
    registry.build_from_financial_sources(sources)

    sbs_long = sbs_loader.load_all_sbs_classifications()
    nombres_unicos = sbs_long[["ENTITY_NAME_RAW", "ENTITY_TYPE_RAW"]].drop_duplicates()
    name_to_id = {}
    for _, row in nombres_unicos.iterrows():
        name_to_id[row["ENTITY_NAME_RAW"]] = registry.resolve_sbs_name(
            row["ENTITY_NAME_RAW"], row["ENTITY_TYPE_RAW"], "ALL"
        )
    sbs_long = sbs_long.copy()
    sbs_long["ENTITY_ID"] = sbs_long["ENTITY_NAME_RAW"].map(name_to_id)
    return sbs_long, registry
